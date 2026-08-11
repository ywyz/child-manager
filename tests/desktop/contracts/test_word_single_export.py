from __future__ import annotations

import errno
import json
import os
from copy import deepcopy
from dataclasses import FrozenInstanceError
from datetime import date
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from typing import Any, cast

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.shared import RGBColor

from kindergarten_manager.application.dto import CancellationToken
from kindergarten_manager.application.exports import (
    DailyPlanExportSnapshot,
    ExportError,
    ExportService,
)
from kindergarten_manager.domain.content import PlanContentV1
from kindergarten_manager.infrastructure.exports.teacherplan_renderer import TeacherplanRenderer
from tests.desktop.helpers import implemented

TEMPLATE_SHA256 = "72ee26e7cb8f510a11bc303b7a967c2a375fe436b5c8a72822ee9ccbfe235043"
FIXTURE = Path("tests/fixtures/word/daily_activity_plan_v1.json")


class FrozenContentSnapshot:
    def __init__(self) -> None:
        self._payload = json.loads(FIXTURE.read_text(encoding="utf-8"))["content_snapshot"]

    def model_dump(self, *, mode: str = "python") -> dict[str, Any]:
        assert mode in {"python", "json"}
        return dict(self._payload)

    def canonical_json(self) -> str:
        return json.dumps(self._payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"))


def _content() -> PlanContentV1:
    return cast(PlanContentV1, FrozenContentSnapshot())


def _snapshot(content: PlanContentV1 | None = None) -> DailyPlanExportSnapshot:
    value = content or _content()
    canonical = value.canonical_json() if hasattr(value, "canonical_json") else "{}"
    return DailyPlanExportSnapshot(
        plan_id=7,
        content_revision=3,
        plan_date=date(2026, 3, 2),
        teaching_week_text="第（四）周",
        activity_date_text="周（一）3月2日",
        semester_name="2026 春季学期",
        semester_start_date=date(2026, 2, 4),
        semester_end_date=date(2026, 6, 30),
        kindergarten_name="星河幼儿园",
        class_name="向日葵班",
        age_group="middle",
        author_name="测试教师",
        content=value,
        content_sha256=sha256(canonical.encode()).hexdigest(),
    )


def _east_asia_font(run: object) -> str | None:
    element = run._element  # type: ignore[attr-defined]
    if element.rPr is None or element.rPr.rFonts is None:
        return None
    return element.rPr.rFonts.get(qn("w:eastAsia"))


def _paragraph_properties_without_numbering(paragraph: object) -> str | None:
    properties = paragraph._p.pPr  # type: ignore[attr-defined]
    if properties is None:
        return None
    normalized = deepcopy(properties)
    if normalized.numPr is not None:
        normalized.remove(normalized.numPr)
    return normalized.xml


def test_snapshot_is_frozen_and_renderer_preserves_template_contract(
    teacherplan_template_path: Path,
) -> None:
    snapshot = _snapshot()
    with pytest.raises(FrozenInstanceError):
        snapshot.content_revision = 4  # type: ignore[misc]
    before = sha256(teacherplan_template_path.read_bytes()).hexdigest()
    renderer = TeacherplanRenderer(teacherplan_template_path, expected_sha256=TEMPLATE_SHA256)

    rendered = implemented(lambda: renderer.render_day(snapshot))
    document = Document(BytesIO(rendered))

    assert sha256(teacherplan_template_path.read_bytes()).hexdigest() == before == TEMPLATE_SHA256
    assert len(document.tables) == 1
    table = document.tables[0]
    assert (len(table.rows), len(table.columns)) == (19, 2)
    assert document.paragraphs[0].text == "星河幼儿园一日活动计划（2026.2-2026.6）"
    assert document.paragraphs[1].text == "向日葵班 测试教师"
    assert table.cell(0, 0).text == "第（四）周"
    assert table.cell(1, 0).text == "周（一）3月2日"
    expected_cells = {
        (2, 1): "体能大循环\n集体游戏：《合作运球》\n自主游戏：《自主器械》",
        (3, 1): (
            "重点指导：合作运球\n活动目标：\n1.练习双人合作。\n2.发展身体协调。"
            "\n3.体验共同游戏。\n指导要点：\n1.保持适当距离。\n2.听清同伴提醒。"
            "\n3.结束后收好器材。"
        ),
        (4, 1): "话题：《春天的发现》",
        (5, 1): "问题设计：\n1.你看到了什么？\n2.你听到了什么？\n3.你想到了什么？",
        (6, 1): "活动主题：《寻找春天》",
        (7, 1): "活动目标：\n1.观察季节变化。\n2.尝试完整表达。",
        (8, 1): "活动准备：\n1.春景图片。\n2.记录纸。",
        (9, 1): "活动重点：表达观察结果。",
        (10, 1): "活动难点：连续描述变化。",
        (11, 1): (
            "活动过程：\n一、观察图片\n1.师：请看看图片里有什么。\n2.幼：我看到了花。"
            "\n二、新增延伸\n1.师：请把发现画下来。"
        ),
        (12, 1): "游戏区域：建构区、阅读区",
        (13, 1): (
            "重点指导：建构区\n活动目标：\n1.尝试合作搭建。\n2.发展空间想象。"
            "\n3.学习整理材料。\n指导要点：\n1.先商量再搭建。\n2.按标记取放材料。"
            "\n3.结束后共同整理。"
        ),
        (14, 1): "支持策略：\n1.提供结构图片。\n2.补充连接材料。\n3.展示幼儿作品。",
        (15, 1): "游戏区域：球类区、平衡区",
        (16, 1): (
            "重点指导：球类区\n活动目标：\n1.练习双手滚球。\n2.发展手眼协调。"
            "\n3.遵守游戏规则。\n指导要点：\n1.面向同伴滚球。\n2.控制滚球力量。"
            "\n3.等待轮流游戏。"
        ),
        (17, 1): "支持策略：\n1.提供不同大小软球。\n2.设置宽窄目标线。\n3.鼓励同伴合作。",
        (18, 1): "活动亮点：\n存在问题：\n调整策略：",
    }
    assert {
        coordinates: table.cell(*coordinates).text.rstrip("\n") for coordinates in expected_cells
    } == expected_cells
    process = table.cell(11, 1)
    red_text = "".join(
        run.text
        for paragraph in process.paragraphs
        for run in paragraph.runs
        if run.font.color.rgb == RGBColor(255, 0, 0)
    )
    non_red_process_text = "".join(
        run.text
        for paragraph in process.paragraphs
        for run in paragraph.runs
        if run.font.color.rgb != RGBColor(255, 0, 0)
    )
    assert "二、新增延伸" in red_text
    assert "请把发现画下来" in red_text
    assert "一、观察图片" in non_red_process_text
    assert "请看看图片里有什么" in non_red_process_text
    assert all(
        run.font.color.rgb != RGBColor(255, 0, 0)
        for row_index in range(len(table.rows))
        for column_index in range(len(table.columns))
        if (row_index, column_index) != (11, 1)
        for paragraph in table.cell(row_index, column_index).paragraphs
        for run in paragraph.runs
    )
    title_run = document.paragraphs[0].runs[0]
    body_run = table.cell(2, 1).paragraphs[0].runs[0]
    assert title_run.font.size is not None and title_run.font.size.pt == 16
    assert title_run.font.name == _east_asia_font(title_run) == "楷体"
    assert body_run.font.size is not None and body_run.font.size.pt == 12
    assert body_run.font.name == _east_asia_font(body_run) == "仿宋"

    template = Document(str(teacherplan_template_path))
    for row_index in range(len(template.tables[0].rows)):
        for column_index in range(len(template.tables[0].columns)):
            before_paragraphs = template.tables[0].cell(row_index, column_index).paragraphs
            after_paragraphs = table.cell(row_index, column_index).paragraphs
            assert len(after_paragraphs) == len(before_paragraphs)
            assert [
                _paragraph_properties_without_numbering(paragraph) for paragraph in after_paragraphs
            ] == [
                _paragraph_properties_without_numbering(paragraph)
                for paragraph in before_paragraphs
            ]


@pytest.mark.parametrize("keep_input_numbers", [False, True])
def test_area_game_lists_render_one_numbering_layer_without_numbered_labels(
    teacherplan_template_path: Path,
    keep_input_numbers: bool,
) -> None:
    def lines(first: str, second: str) -> list[str]:
        if keep_input_numbers:
            return [f"1.{first}", f"2.{second}"]
        return [first, second]

    payload = PlanContentV1.empty().model_dump()
    payload["indoor_area_game"] = {
        "areas": ["娃娃家", "阅读区"],
        "focus_guidance": "娃娃家",
        "objectives": lines("娃娃家", "阅读区"),
        "guidance_points": lines("娃娃家指导要点", "阅读区指导要点"),
        "support_strategies": lines("娃娃家支持策略", "阅读区支持策略"),
    }
    payload["afternoon_outdoor_game"] = {
        "areas": ["七星区"],
        "focus_guidance": "玩水区",
        "objectives": lines("欸哦赛", "alksdjnk"),
        "guidance_points": lines("七星区指导要点", "阅读区指导要点"),
        "support_strategies": lines("骑行区支持策略", "阅读区支持策略"),
    }
    renderer = TeacherplanRenderer(teacherplan_template_path, expected_sha256=TEMPLATE_SHA256)

    rendered = renderer.render_day(_snapshot(PlanContentV1.model_validate(payload)))
    table = Document(BytesIO(rendered)).tables[0]

    expected = {
        13: [
            "重点指导：娃娃家",
            "活动目标：",
            "1.娃娃家",
            "2.阅读区",
            "指导要点：",
            "1.娃娃家指导要点",
            "2.阅读区指导要点",
        ],
        14: ["支持策略：", "1.娃娃家支持策略", "2.阅读区支持策略"],
        16: [
            "重点指导：玩水区",
            "活动目标：",
            "1.欸哦赛",
            "2.alksdjnk",
            "指导要点：",
            "1.七星区指导要点",
            "2.阅读区指导要点",
        ],
        17: ["支持策略：", "1.骑行区支持策略", "2.阅读区支持策略"],
    }
    for row_index, expected_lines in expected.items():
        paragraphs = table.cell(row_index, 1).paragraphs
        assert [paragraph.text for paragraph in paragraphs if paragraph.text] == expected_lines
        assert all(
            paragraph._p.pPr is None or paragraph._p.pPr.numPr is None for paragraph in paragraphs
        )


class BytesRenderer:
    def __init__(self, payload: bytes | None = None, failure: Exception | None = None) -> None:
        if payload is None:
            document = Document("templates/teacherplan/teacherplan.docx")
            document.tables[0].cell(0, 0).text = "第（四）周"
            document.tables[0].cell(1, 0).text = "周（一）3月2日"
            stream = BytesIO()
            document.save(stream)
            payload = stream.getvalue()
        self.payload = payload
        self.failure = failure

    def render_day(self, _snapshot: DailyPlanExportSnapshot) -> bytes:
        if self.failure is not None:
            raise self.failure
        return self.payload

    def validate_day(self, path: Path, snapshot: DailyPlanExportSnapshot) -> None:
        document = Document(str(path))
        if len(document.tables) != 1:
            raise ValueError("表格数量不匹配")
        table = document.tables[0]
        if (len(table.rows), len(table.columns)) != (19, 2):
            raise ValueError("模板表格结构不匹配")
        if table.cell(0, 0).text != snapshot.teaching_week_text:
            raise ValueError("教学周不匹配")
        if table.cell(1, 0).text != snapshot.activity_date_text:
            raise ValueError("活动日期不匹配")


def _opaque_snapshot() -> DailyPlanExportSnapshot:
    return _snapshot(cast(PlanContentV1, object()))


class SnapshotReader:
    def __init__(self, snapshot: DailyPlanExportSnapshot | None) -> None:
        self.snapshot = snapshot
        self.requested_plan_ids: list[int] = []

    def load_export_snapshot(self, plan_id: int) -> DailyPlanExportSnapshot | None:
        self.requested_plan_ids.append(plan_id)
        return self.snapshot


def test_prepare_single_returns_frozen_snapshot_and_rejects_missing_plan() -> None:
    snapshot = _snapshot()
    reader = SnapshotReader(snapshot)
    service = ExportService(BytesRenderer(), snapshot_reader=reader)

    prepared = implemented(lambda: service.prepare_single(7))

    assert prepared == snapshot
    assert reader.requested_plan_ids == [7]
    with pytest.raises(ExportError) as captured:
        implemented(
            lambda: ExportService(
                BytesRenderer(), snapshot_reader=SnapshotReader(None)
            ).prepare_single(404)
        )
    assert captured.value.error_code == "plan.not_found"


def test_single_export_uses_same_directory_temporary_file_and_atomic_replace(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    destination = tmp_path / "一日活动计划.docx"
    destination.write_bytes(b"old-docx")
    renderer = BytesRenderer()
    service = ExportService(renderer)
    fsync_calls: list[int] = []
    real_fsync = os.fsync

    def recording_fsync(file_descriptor: int) -> None:
        fsync_calls.append(file_descriptor)
        real_fsync(file_descriptor)

    monkeypatch.setattr(os, "fsync", recording_fsync)

    result = implemented(
        lambda: service.export_single(_opaque_snapshot(), destination, CancellationToken())
    )

    assert result.destination == destination
    assert result.exported_dates == (date(2026, 3, 2),)
    assert result.skipped == ()
    assert destination.read_bytes() == renderer.payload
    reopened = Document(str(destination))
    assert len(reopened.tables) == 1
    assert reopened.tables[0].cell(1, 0).text == "周（一）3月2日"
    assert fsync_calls
    assert list(tmp_path.iterdir()) == [destination]


def test_cancel_or_render_failure_preserves_existing_destination(tmp_path: Path) -> None:
    destination = tmp_path / "一日活动计划.docx"
    destination.write_bytes(b"old-docx")
    cancelled = CancellationToken()
    cancelled.request_cancel()

    with pytest.raises(ExportError) as captured:
        implemented(
            lambda: ExportService(BytesRenderer()).export_single(
                _opaque_snapshot(), destination, cancelled
            )
        )
    assert captured.value.error_code == "export.cancelled"
    assert destination.read_bytes() == b"old-docx"

    with pytest.raises(ExportError) as captured:
        implemented(
            lambda: ExportService(
                BytesRenderer(failure=OSError("synthetic disk failure"))
            ).export_single(_opaque_snapshot(), destination, CancellationToken())
        )
    assert captured.value.error_code == "export.render_failed"
    assert destination.read_bytes() == b"old-docx"
    assert list(tmp_path.iterdir()) == [destination]


def test_no_space_or_locked_destination_has_stable_error_and_cleans_temporary_file(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    destination = tmp_path / "一日活动计划.docx"
    destination.write_bytes(b"old-docx")

    with pytest.raises(ExportError) as captured:
        implemented(
            lambda: ExportService(
                BytesRenderer(failure=OSError(errno.ENOSPC, "synthetic disk full"))
            ).export_single(_opaque_snapshot(), destination, CancellationToken())
        )
    assert captured.value.error_code == "export.no_space"

    def locked(_source: os.PathLike[str], _destination: os.PathLike[str]) -> None:
        raise PermissionError("synthetic Word lock")

    monkeypatch.setattr(os, "replace", locked)
    with pytest.raises(ExportError) as captured:
        implemented(
            lambda: ExportService(BytesRenderer()).export_single(
                _opaque_snapshot(), destination, CancellationToken()
            )
        )
    assert captured.value.error_code == "export.destination_locked"
    assert destination.read_bytes() == b"old-docx"
    assert list(tmp_path.iterdir()) == [destination]

    def publish_failed(_source: object, _destination: object) -> None:
        raise OSError(errno.EIO, "synthetic publish failure")

    monkeypatch.setattr(os, "replace", publish_failed)
    with pytest.raises(ExportError) as captured:
        implemented(
            lambda: ExportService(BytesRenderer()).export_single(
                _opaque_snapshot(), destination, CancellationToken()
            )
        )
    assert captured.value.error_code == "export.publish_failed"
    assert destination.read_bytes() == b"old-docx"
    assert list(tmp_path.iterdir()) == [destination]


def test_invalid_rendered_docx_is_rejected_before_publish(tmp_path: Path) -> None:
    destination = tmp_path / "一日活动计划.docx"
    destination.write_bytes(b"old-docx")
    invalid = Document()
    invalid.add_paragraph("缺少模板表格与日期")
    stream = BytesIO()
    invalid.save(stream)

    with pytest.raises(ExportError) as captured:
        implemented(
            lambda: ExportService(BytesRenderer(payload=stream.getvalue())).export_single(
                _opaque_snapshot(), destination, CancellationToken()
            )
        )
    assert captured.value.error_code == "export.validation_failed"
    assert destination.read_bytes() == b"old-docx"
    assert list(tmp_path.iterdir()) == [destination]
