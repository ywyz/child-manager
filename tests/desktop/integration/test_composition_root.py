from __future__ import annotations

from datetime import date
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from PySide6.QtCore import QDate, Qt
from PySide6.QtWidgets import (
    QComboBox,
    QDateEdit,
    QFileDialog,
    QLineEdit,
    QPlainTextEdit,
    QPushButton,
    QStackedWidget,
    QWidget,
)
from pytestqt.qtbot import QtBot

from kindergarten_manager.app import create_desktop_window
from kindergarten_manager.infrastructure.paths import DesktopPaths
from tests.desktop.helpers import implemented


def _child(window: QWidget, widget_type: type[Any], name: str) -> Any:
    child = window.findChild(widget_type, name)
    assert child is not None, name
    return child


def test_composition_root_persists_first_daily_plan_across_restart_and_exports_word(
    qtbot: QtBot,
    tmp_path: Path,
    teacherplan_template_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    root = tmp_path / "GenericDataLocation" / "cn.kindergartenmanager.desktop"
    paths = DesktopPaths(
        root=root,
        data=root / "data",
        database=root / "data" / "child-manager.sqlite3",
        backups=root / "backups",
        daily_backups=root / "backups" / "daily",
        pre_migration_backups=root / "backups" / "pre-migration",
        pre_restore_backups=root / "backups" / "pre-restore",
        recovery=root / "recovery",
        staging=root / "staging",
        cache=root / "cache",
        logs=root / "logs",
    )

    window = implemented(
        lambda: create_desktop_window(
            paths=paths,
            template_path=teacherplan_template_path,
            today=lambda: date(2026, 9, 7),
        )
    )
    qtbot.addWidget(window)
    assert window.windowTitle() == "幼儿园管理助手"

    stack = window.findChild(QStackedWidget, "main_stack")
    assert paths.database.exists()
    assert stack is not None
    assert stack.currentIndex() == 0

    values = {
        "teacher_name": "测试教师",
        "kindergarten_name": "星河幼儿园",
        "semester_name": "2026 秋季学期",
        "class_name": "向日葵班",
    }
    for object_name, value in values.items():
        field = _child(window, QLineEdit, object_name)
        assert isinstance(field, QLineEdit)
        field.setText(value)
    _child(window, QDateEdit, "semester_start_date").setDate(QDate(2026, 9, 1))
    _child(window, QDateEdit, "semester_end_date").setDate(QDate(2027, 1, 31))
    complete = _child(window, QPushButton, "complete_setup")
    qtbot.mouseClick(complete, Qt.MouseButton.LeftButton)
    qtbot.mouseClick(_child(window, QPushButton, "open_settings"), Qt.MouseButton.LeftButton)
    theme_preference = _child(window, QComboBox, "theme_preference")
    theme_preference.setCurrentIndex(theme_preference.findData("dark"))
    _child(window, QLineEdit, "semester_settings_name").setText("2026—2027 学年")
    _child(window, QDateEdit, "semester_settings_start_date").setDate(QDate(2026, 8, 20))
    _child(window, QDateEdit, "semester_settings_end_date").setDate(QDate(2027, 7, 15))
    _child(window, QLineEdit, "kindergarten_settings_name").setText("星河实验幼儿园")
    _child(window, QLineEdit, "teacher_settings_name").setText("测试教师二")
    _child(window, QLineEdit, "class_settings_name").setText("向日葵实验班")
    age_group = _child(window, QComboBox, "class_settings_age_group")
    age_group.setCurrentIndex(age_group.findData("large"))
    _child(window, QPlainTextEdit, "class_settings_indoor_areas").setPlainText("建构区\n美工区")
    _child(window, QPlainTextEdit, "class_settings_outdoor_areas").setPlainText("沙池\n操场")
    qtbot.mouseClick(_child(window, QPushButton, "save_settings"), Qt.MouseButton.LeftButton)
    qtbot.mouseClick(_child(window, QPushButton, "back_to_plan"), Qt.MouseButton.LeftButton)
    theme = _child(window, QLineEdit, "group_activity_theme")
    source = _child(window, QPlainTextEdit, "group_activity_source_text")
    assert isinstance(theme, QLineEdit)
    source.setPlainText("活动名称：寻找秋天\n活动过程：观察落叶。")
    theme.setText("寻找秋天")
    qtbot.mouseClick(
        _child(window, QPushButton, "save_plan"),
        Qt.MouseButton.LeftButton,
    )
    date_context = _child(window, QDateEdit, "plan_date")
    original_date = date_context.date()
    date_context.setDate(original_date.addDays(1))
    assert theme.text() == ""
    theme.setText("第二天的活动")
    qtbot.mouseClick(_child(window, QPushButton, "save_plan"), Qt.MouseButton.LeftButton)
    date_context.setDate(original_date)
    assert theme.text() == "寻找秋天"
    window.close()

    reopened = implemented(
        lambda: create_desktop_window(
            paths=paths,
            template_path=teacherplan_template_path,
            today=lambda: date(2026, 9, 7),
        )
    )
    qtbot.addWidget(reopened)
    reopened_theme = _child(reopened, QLineEdit, "group_activity_theme")
    assert isinstance(reopened_theme, QLineEdit)
    assert reopened_theme.text() == "寻找秋天"
    assert (
        _child(reopened, QPlainTextEdit, "group_activity_source_text").toPlainText()
        == "活动名称：寻找秋天\n活动过程：观察落叶。"
    )
    qtbot.mouseClick(_child(reopened, QPushButton, "open_settings"), Qt.MouseButton.LeftButton)
    reopened_preference = _child(reopened, QComboBox, "theme_preference")
    assert reopened_preference.currentData() == "dark"
    assert _child(reopened, QLineEdit, "semester_settings_name").text() == "2026—2027 学年"
    assert _child(reopened, QDateEdit, "semester_settings_start_date").date() == QDate(2026, 8, 20)
    assert _child(reopened, QDateEdit, "semester_settings_end_date").date() == QDate(2027, 7, 15)
    assert _child(reopened, QLineEdit, "kindergarten_settings_name").text() == "星河实验幼儿园"
    assert _child(reopened, QLineEdit, "teacher_settings_name").text() == "测试教师二"
    assert _child(reopened, QLineEdit, "class_settings_name").text() == "向日葵实验班"
    assert _child(reopened, QComboBox, "class_settings_age_group").currentData() == "large"
    qtbot.mouseClick(_child(reopened, QPushButton, "back_to_plan"), Qt.MouseButton.LeftButton)

    destination = tmp_path / "当天教案.docx"

    def choose_destination(*args: object, **_kwargs: object) -> tuple[str, str]:
        assert "一日活动计划-向日葵实验班-2026-09-07.docx" in str(args[2])
        return str(destination), "Word 文档 (*.docx)"

    monkeypatch.setattr(
        QFileDialog,
        "getSaveFileName",
        choose_destination,
    )
    qtbot.mouseClick(
        _child(reopened, QPushButton, "export_day"),
        Qt.MouseButton.LeftButton,
    )
    assert destination.is_file()
    exported = Document(str(destination))
    assert exported.paragraphs[0].text == "星河实验幼儿园一日活动计划（2026.8-2027.7）"
    assert exported.tables[0].cell(6, 1).text == "活动主题：《寻找秋天》"
